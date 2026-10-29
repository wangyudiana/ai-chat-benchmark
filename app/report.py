# -*- coding: utf-8 -*-
"""截图对照 Excel 和 HTML 调研报告。

报告 = 系统自动生成的数据部分（评测设计、排名、分维度/分类、速度、典型案例截图，保证数字准确）
     + AI 按内置分析框架写的结论部分（见 analyzer.report_prompt）。"""
import base64
import html
import io
import re

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image

from . import store
from .analyzer import compute_stats, f1, fs, item_avg, pick_questions

FONT = "Arial"
MAX_SHOTS = 800          # 截图对照里最多放这么多张，再多文件会太大
THUMB_W, THUMB_H = 420, 540     # 在 Excel 格子里显示的大小（像素）
THUMB_SCALE = 3                  # 图片本身按 3 倍像素存，放大看、双击看原图都清楚


def thumb(path, width=THUMB_W, max_h=THUMB_H, quality=88, scale=THUMB_SCALE):
    """按显示宽度的 scale 倍缩放（原图更小就不放大）；长图只保留开头一段（完整长图另给链接）。
    返回 (JPEG 字节, 是否截断, 显示高度)。"""
    im = Image.open(path).convert("RGB")
    pw = min(im.width, width * scale)
    ph = int(im.height * pw / im.width)
    if pw != im.width:
        im = im.resize((pw, max(1, ph)), Image.LANCZOS)
    k = pw / width                       # 实际像素 / 显示像素
    cut = ph > max_h * k
    if cut:
        im = im.crop((0, 0, pw, int(max_h * k)))
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=quality, optimize=True, subsampling=0)
    return buf.getvalue(), cut, round(im.height / k)


def opt_desc(cfg, key, sites_all):
    opts = (cfg.get("site_opts") or {}).get(key) or {}
    labels = {o["id"]: o["label"] for o in (sites_all.get(key) or {}).get("options", [])}
    return "、".join(f"{labels.get(k, k)}：{'开' if v is True else '关' if v is False else v}" for k, v in opts.items()) or "默认"


# ---------------------------------------------------------------- 截图对照 Excel
def build_shots_xlsx(jid):
    job, res = store.get_job(jid), store.get_results(jid)
    base = store.job_dir(jid)
    sites = res["sites"]
    by = {(r["query_id"], r["site"]): r for r in res["results"]}
    wb = Workbook()
    ws = wb.active
    ws.title = "截图对照"
    ws.append([f"{job['name']}　截图对照（每格是长截图的开头部分；点下面一行的链接看完整长图，需要在“打包下载”解压后的文件夹里打开）"])
    ws.append(["编号", "分类", "问题", *[s["name"] for s in sites]])
    for c in ws[2]:
        c.font = Font(name=FONT, bold=True)
        c.fill = PatternFill("solid", fgColor="E3E8FB")
    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 34
    for i in range(len(sites)):
        ws.column_dimensions[get_column_letter(4 + i)].width = 60
    ws.freeze_panes = "D3"
    n_img = 0
    row = 3
    for q in res["queries"]:
        if not any(by.get((q["id"], s["key"])) for s in sites):
            continue
        ws.cell(row, 1, q["id"])
        ws.cell(row, 2, q["category"])
        ws.cell(row, 3, " → ".join(q.get("turns") or [q["query"]]))
        ws.row_dimensions[row].height = THUMB_H * 0.75 + 6
        for i, s in enumerate(sites):
            col = 4 + i
            r = by.get((q["id"], s["key"]))
            link = ws.cell(row + 1, col)
            if not r:
                ws.cell(row, col, "未采集")
                continue
            shot = r.get("screenshot") or r.get("window_screenshot")
            meta = [{"ok": "成功", "timeout": "超时", "error": "失败"}.get(r["status"], r["status"])]
            if r.get("total_ms") is not None:
                meta.append(f"完成 {r['total_ms'] / 1000:.1f}s")
            meta.append(f"{len(r.get('response_text') or '')} 字")
            if shot and (base / shot).exists() and n_img < MAX_SHOTS:
                try:
                    data, cut, disp_h = thumb(base / shot)
                    img = XLImage(io.BytesIO(data))
                    img.width, img.height = THUMB_W, disp_h      # 显示尺寸；图片本身是 3 倍像素
                    ws.add_image(img, f"{get_column_letter(col)}{row}")
                    n_img += 1
                    meta.append("长图已截取开头" if cut else "")
                except Exception:
                    ws.cell(row, col, "截图读取失败")
                link.value = "完整长图　" + "　".join(m for m in meta if m)
                link.hyperlink = shot
                link.font = Font(name=FONT, color="1F4FD8", underline="single", size=9)
            else:
                ws.cell(row, col, (r.get("error") or "没有截图") if n_img < MAX_SHOTS else f"截图超过 {MAX_SHOTS} 张，其余见打包下载")
                link.value = "　".join(meta)
                link.font = Font(name=FONT, size=9, color="666666")
        for c in ws[row]:
            if c.column <= 3:
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.font = Font(name=FONT)
        row += 2
    ws["A1"].font = Font(name=FONT, bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- Markdown -> HTML（只支持报告里会用到的写法）
def md_to_html(md):
    out, para, lst, code, table = [], [], None, None, []

    def inline(t):
        t = html.escape(t)
        t = re.sub(r"`([^`]+)`", r"<code>\1</code>", t)
        t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
        t = re.sub(r"(?<!\*)\*([^*\n]+)\*(?!\*)", r"<i>\1</i>", t)
        t = re.sub(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", r'<a href="\2" target="_blank">\1</a>', t)
        return t

    def flush():
        nonlocal para, lst, table
        if para:
            out.append("<p>" + "<br>".join(inline(x) for x in para) + "</p>")
            para = []
        if lst:
            tag, items = lst
            out.append(f"<{tag}>" + "".join(f"<li>{inline(x)}</li>" for x in items) + f"</{tag}>")
            lst = None
        if table:
            rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in table if not re.fullmatch(r"\s*\|?[\s:|-]+\|?\s*", r)]
            if rows:
                out.append("<table><thead><tr>" + "".join(f"<th>{inline(c)}</th>" for c in rows[0]) + "</tr></thead><tbody>"
                           + "".join("<tr>" + "".join(f"<td>{inline(c)}</td>" for c in r) + "</tr>" for r in rows[1:]) + "</tbody></table>")
            table = []

    for line in (md or "").splitlines():
        if code is not None:
            if line.strip().startswith("```"):
                out.append("<pre><code>" + html.escape("\n".join(code)) + "</code></pre>")
                code = None
            else:
                code.append(line)
            continue
        s = line.strip()
        if s.startswith("```"):
            flush()
            code = []
        elif not s:
            flush()
        elif m := re.match(r"(#{1,6})\s+(.*)", s):
            flush()
            lv = min(len(m.group(1)) + 1, 6)
            out.append(f"<h{lv}>{inline(m.group(2))}</h{lv}>")
        elif s.startswith("|"):
            if para or lst:
                flush()
            table.append(s)
        elif m := re.match(r"([-*•]|\d+[.、)])\s+(.*)", s):
            if para or table:
                flush()
            tag = "ol" if m.group(1)[0].isdigit() else "ul"
            if lst and lst[0] != tag:
                flush()
            lst = lst or (tag, [])
            lst[1].append(m.group(2))
        elif s.startswith(">"):
            flush()
            out.append(f"<blockquote>{inline(s.lstrip('> '))}</blockquote>")
        else:
            if lst or table:
                flush()
            para.append(s)
    flush()
    if code is not None:
        out.append("<pre><code>" + html.escape("\n".join(code)) + "</code></pre>")
    return "\n".join(out)


# ---------------------------------------------------------------- HTML 报告
CSS = """
body{font:15px/1.7 -apple-system,"PingFang SC","Microsoft YaHei",Arial,sans-serif;color:#1d2330;background:#f4f6f9;margin:0}
.wrap{max-width:1080px;margin:0 auto;padding:36px 28px 80px}
h1{font-size:26px;margin:0 0 6px}h2{font-size:19px;margin:40px 0 12px;padding-top:8px;border-top:1px solid #dde2ea}
h3{font-size:16px;margin:22px 0 8px}h4{font-size:15px;margin:16px 0 6px}
.meta{color:#5b6577;font-size:13px}.meta span{margin-right:14px}
.card{background:#fff;border:1px solid #dde2ea;border-radius:10px;padding:16px 20px;margin:12px 0}
.ai{border-left:4px solid #3b5bdb}
table{border-collapse:collapse;width:100%;font-size:14px;background:#fff}
th,td{border-bottom:1px solid #e6e9ef;padding:8px 10px;text-align:left;vertical-align:top}th{color:#5b6577;font-weight:500;background:#f8f9fb}
td.n{text-align:right;font-variant-numeric:tabular-nums}
.bar{display:flex;align-items:center;gap:10px;margin:6px 0}.bar .nm{width:120px}.bar .tr{flex:1;background:#eef1f5;border-radius:5px;height:14px}
.bar .tr i{display:block;height:100%;border-radius:5px;background:#3b5bdb}.bar .v{width:44px;text-align:right;font-weight:600}
.case{display:grid;grid-template-columns:repeat(auto-fill,minmax(240px,1fr));gap:12px}
.case .it{background:#fff;border:1px solid #dde2ea;border-radius:8px;padding:10px}
.case img{width:100%;border:1px solid #eef1f5;border-radius:4px;margin-top:6px;cursor:zoom-in}
.case img.big{position:fixed;inset:2vh 2vw;width:96vw;height:96vh;object-fit:contain;background:rgba(20,24,32,.92);z-index:9;cursor:zoom-out;margin:0;border:0}
.small{font-size:12.5px;color:#5b6577}.best{background:#ffe04f;border-radius:3px;padding:0 5px;font-size:12px;margin-left:6px}
code{background:#eef1f5;border-radius:4px;padding:0 4px}pre{background:#f3f5f8;padding:10px;border-radius:6px;overflow:auto}
details{background:#fff;border:1px solid #dde2ea;border-radius:8px;padding:8px 14px;margin:8px 0}summary{cursor:pointer}
blockquote{margin:8px 0;padding:4px 12px;border-left:3px solid #ccd3df;color:#5b6577}
"""


def pct(v):
    return "—" if v is None else f"{v * 100:.0f}%"


def heat(v):
    if v is None:
        return '<td class="n">—</td>'
    a = max(0.0, min(1.0, (v - 1) / 4))
    return f'<td class="n" style="background:rgba(59,91,219,{a * 0.35:.2f})">{v:.1f}</td>'


def build_html(jid, with_images=True):
    job, res, an = store.get_job(jid), store.get_results(jid), store.get_analysis(jid)
    cfg = job["config"]
    base = store.job_dir(jid)
    st, cats = compute_stats(res, an)
    dims = an["dims"]
    single = len(res["sites"]) == 1
    sites_all = {s["key"]: s for s in store.get_sites()}
    by = {(r["query_id"], r["site"]): r for r in res["results"]}
    names = {s["key"]: s["name"] for s in res["sites"]}
    e = html.escape
    ranked = sorted(st, key=lambda s: -(s["overall"] or -1))
    scored = any(s["scored"] for s in st)
    n_q = len({q.get("base_id") or q["id"] for q in res["queries"]})
    H = [f"<!doctype html><html lang='zh-CN'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
         f"<title>{e(job['name'])}</title><style>{CSS}</style></head><body><div class='wrap'>",
         f"<h1>{e(job['name'])}</h1><div class='meta'><span>创建 {e(job['created_at'][:16].replace('T', ' '))}</span>"
         f"<span>{len(res['sites'])} 个产品</span><span>{n_q} 题</span>"
         f"<span>{'未登录浏览器' if cfg.get('profile') == 'guest' else '无痕对话' if cfg.get('incognito') else '普通对话'}</span>"
         f"<span>{'手机网页' if cfg.get('device') == 'mobile' else '电脑网页'}</span></div>"]
    # 评测设计
    H.append("<h2>评测设计</h2><div class='card'>")
    if an.get("purpose"):
        H.append(f"<p><b>评测目的：</b>{e(an['purpose'])}</p>")
    H.append("<p><b>产品与模型/模式：</b>" + "；".join(f"{e(s['name'])}（{e(opt_desc(cfg, s['key'], sites_all))}）" for s in res["sites"]) + "</p>")
    H.append(f"<p><b>问题：</b>{n_q} 题" + (f"，每题跑 {cfg['repeat']} 次" if cfg.get("repeat", 1) > 1 else "")
             + (f"；分类：{e('、'.join(cats))}" if len(cats) > 1 else "") + "</p>")
    desc = an.get("dim_desc") or {}
    H.append("<p><b>评分维度（1–5 分）：</b></p><ul>" + "".join(f"<li>{e(d)}" + (f"：{e(desc[d])}" if desc.get(d) else "") + "</li>" for d in dims) + "</ul>")
    n_ai = sum(1 for it in an["items"].values() if it.get("src") == "ai" and item_avg(it, dims) is not None)
    n_hu = sum(1 for it in an["items"].values() if it.get("src") != "ai" and item_avg(it, dims) is not None)
    H.append(f"<p class='small'>评分方式：AI 评审 {n_ai} 条、人工 {n_hu} 条（人工优先）。AI 评审时隐去产品名并打乱顺序，减少品牌和位置偏好。"
             "首字/完成时间在浏览器里测得，适合横向比较快慢。</p></div>")
    # AI 结论
    H.append(f"<h2>{'评测结论' if single else '竞品结论'}</h2>")
    H.append(f"<div class='card ai'>{md_to_html(an['report'])}</div>" if an.get("report")
             else "<div class='card small'>还没有生成结论：到平台的“分析”页点“让 AI 写结论”。</div>")
    # 数据
    if scored:
        top = max((s["overall"] or 0) for s in st) or 5
        H.append(f"<h2>{'综合得分' if single else '综合排名'}</h2><div class='card'>" + "".join(
            f"<div class='bar'><span class='nm'>{e(s['name'])}</span><span class='tr'><i style='width:{(s['overall'] or 0) / 5 * 100:.0f}%'></i></span>"
            f"<span class='v'>{f1(s['overall'])}</span><span class='small'>已评 {s['scored']}/{s['n']}" + ("" if single else f"，最佳 {s['wins']} 次") + "</span></div>"
            for s in ranked) + "</div>")
        H.append("<h3>分维度</h3><table><thead><tr><th>产品</th>" + "".join(f"<th class='n'>{e(d)}</th>" for d in dims)
                 + "<th class='n'>综合</th></tr></thead><tbody>" + "".join(
                     f"<tr><td>{e(s['name'])}</td>" + "".join(heat(s["dim"][d]) for d in dims) + heat(s["overall"]) + "</tr>" for s in ranked)
                 + "</tbody></table>")
        if len(cats) > 1:
            H.append("<h3>分类表现</h3><table><thead><tr><th>产品</th>" + "".join(f"<th class='n'>{e(c)}</th>" for c in cats)
                     + "</tr></thead><tbody>" + "".join(
                         f"<tr><td>{e(s['name'])}</td>" + "".join(heat(s["cat"][c]) for c in cats) + "</tr>" for s in ranked) + "</tbody></table>")
    H.append("<h3>采集与速度</h3><table><thead><tr><th>产品</th><th class='n'>成功率</th><th class='n'>平均首字</th>"
             "<th class='n'>平均完成</th><th class='n'>平均字数</th></tr></thead><tbody>" + "".join(
                 f"<tr><td>{e(s['name'])}</td><td class='n'>{pct(s['ok_rate'])}</td>"
                 f"<td class='n'>{fs(s['ttft'])}</td><td class='n'>{fs(s['total'])}</td><td class='n'>{'—' if s['chars'] is None else round(s['chars'])}</td></tr>"
                 for s in st) + "</tbody></table>")
    # 典型案例
    picks = pick_questions(res, an, 4)
    if picks:
        H.append(f"<h2>{'得分最低的问题' if single else '典型案例：各家差距最大的问题'}</h2>")
        for p in picks:
            q = p["q"]
            ai = an["ai"].get(q["id"]) or {}
            H.append(f"<h3>{e(q['id'])}　{e(' → '.join(q.get('turns') or [q['query']]))}</h3>")
            if ai.get("summary") and not single:
                H.append(f"<p class='small'>AI 对比：{e(ai['summary'])}</p>")
            H.append("<div class='case'>")
            for s in res["sites"]:
                r = by.get((q["id"], s["key"]))
                it = an["items"].get(f"{q['id']}|{s['key']}") or {}
                av = item_avg(it, dims)
                img = ""
                shot = r and (r.get("screenshot") or r.get("window_screenshot"))
                if with_images and shot and (base / shot).exists():
                    try:
                        data, _, _ = thumb(base / shot, 420, 900, 80, 2)
                        img = (f"<img alt='{e(s['name'])} 截图' title='点击放大' onclick=\"this.classList.toggle('big')\" "
                               f"src='data:image/jpeg;base64,{base64.b64encode(data).decode()}'>")
                    except Exception:
                        pass
                best = "<span class='best'>最佳</span>" if an["best"].get(q["id"]) == s["key"] else ""
                com = (ai.get("comments") or {}).get(s["key"]) or ""
                H.append(f"<div class='it'><b>{e(s['name'])}</b>{best}<span class='small'>　均分 {f1(av)}</span>"
                         + (f"<div class='small'>{e(com)}</div>" if com else "")
                         + (f"<div class='small'>备注：{e(it['note'])}</div>" if it.get("note") else "") + img + "</div>")
            H.append("</div>")
    # 逐条记录
    H.append("<h2>逐条记录</h2>")
    for q in res["queries"]:
        rows = []
        for s in res["sites"]:
            r = by.get((q["id"], s["key"]))
            it = an["items"].get(f"{q['id']}|{s['key']}") or {}
            com = ((an["ai"].get(q["id"]) or {}).get("comments") or {}).get(s["key"]) or ""
            st_txt = "未采集" if not r else {"ok": "", "timeout": "超时", "error": "失败"}.get(r["status"], r["status"])
            rows.append(f"<tr><td>{e(s['name'])}{'<span class=best>最佳</span>' if an['best'].get(q['id']) == s['key'] else ''}</td>"
                        f"<td class='n'>{f1(item_avg(it, dims))}</td><td>{e(st_txt)} {e('、'.join(it.get('tags') or []))}</td>"
                        f"<td>{e(com)}{('<br>备注：' + e(it['note'])) if it.get('note') else ''}</td></tr>")
        H.append(f"<details><summary>{e(q['id'])}　{e(q['category'])}　{e(q['query'][:60])}</summary>"
                 f"<table><thead><tr><th>产品</th><th class='n'>均分</th><th>状态/标签</th><th>点评</th></tr></thead><tbody>{''.join(rows)}</tbody></table></details>")
    H.append("</div></body></html>")
    return "\n".join(H)
